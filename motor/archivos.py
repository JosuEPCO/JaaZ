"""Lectura y escritura de libros: Excel (.xlsx, .xlsm, .xls), CSV y texto (.txt, .ped, .dat).

Los archivos de texto siguen las reglas de Pedigree Viewer: separador coma, punto y
coma, tabulación o espacios (varios espacios cuentan como uno), con comillas para
textos con espacios.
"""
import csv
import io
import math
import re
from datetime import date, datetime, time

LIMITE_CELDAS = 12_000_000
EXT_EXCEL = {".xlsx", ".xlsm", ".xltx", ".xltm"}
EXT_TEXTO = {".csv", ".txt", ".ped", ".dat", ".tsv", ".prn"}

_NUM = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_NUM_COMA = re.compile(r"^[+-]?\d+,\d+$")
_MILES_PUNTO = re.compile(r"^[+-]?\d{1,3}(\.\d{3})+(,\d+)?$")
# "007", "0012": códigos (aretes, registros) cuyo cero inicial es parte del identificador
_CERO_INICIAL = re.compile(r"^[+-]?0\d")
# caracteres de control que Excel no admite dentro de una celda
_ILEGALES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ErrorArchivo(ValueError):
    pass


def _valor_celda(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, (int,)):
        return v
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
        return v
    if isinstance(v, datetime):
        if v.time() == time(0, 0):
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    s = str(v).strip()
    return s if s != "" else None


def _recortar(filas):
    """Quita filas y columnas vacías al final."""
    while filas and all(c is None for c in filas[-1]):
        filas.pop()
    ancho = 0
    for f in filas:
        for j in range(len(f) - 1, -1, -1):
            if f[j] is not None:
                ancho = max(ancho, j + 1)
                break
    return [list(f[:ancho]) + [None] * (ancho - len(f[:ancho])) for f in filas]


def _leer_xlsx(datos):
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:  # archivo dañado o protegido
        raise ErrorArchivo(f"No se pudo abrir el libro de Excel: {e}")
    hojas = []
    celdas = 0
    for ws in libro.worksheets:
        filas = []
        for fila in ws.iter_rows(values_only=True):
            filas.append([_valor_celda(v) for v in fila])
            celdas += len(fila)
            if celdas > LIMITE_CELDAS:
                raise ErrorArchivo("El libro es demasiado grande (más de 12 millones de celdas).")
        hojas.append({"nombre": ws.title, "filas": _recortar(filas)})
    libro.close()
    return hojas


def _leer_xls(datos):
    import xlrd
    try:
        libro = xlrd.open_workbook(file_contents=datos)
    except Exception as e:
        raise ErrorArchivo(f"No se pudo abrir el libro de Excel 97-2003: {e}")
    hojas = []
    for sh in libro.sheets():
        filas = []
        for i in range(sh.nrows):
            fila = []
            for j in range(sh.ncols):
                c = sh.cell(i, j)
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        fila.append(_valor_celda(xlrd.xldate_as_datetime(c.value, libro.datemode)))
                    except Exception:
                        fila.append(c.value)
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    fila.append(None)
                elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                    fila.append(_valor_celda(bool(c.value)))
                elif c.ctype == xlrd.XL_CELL_ERROR:
                    fila.append(None)
                else:
                    fila.append(_valor_celda(c.value))
            filas.append(fila)
        hojas.append({"nombre": sh.name, "filas": _recortar(filas)})
    return hojas


def _decodificar(datos):
    for cod in ("utf-8-sig", "utf-16"):
        try:
            texto = datos.decode(cod)
            if cod == "utf-16" and not datos[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return texto
        except UnicodeDecodeError:
            continue
    return datos.decode("cp1252", errors="replace")


def detectar_separador(texto):
    lineas = [l for l in texto.splitlines()[:30] if l.strip()]
    if not lineas:
        return ","
    candidatos = {}
    for sep in ("\t", ";", ",", "|"):
        cuentas = [l.count(sep) for l in lineas]
        if cuentas[0] > 0 and len(set(cuentas[:min(10, len(cuentas))])) <= 2:
            candidatos[sep] = cuentas[0]
    if candidatos:
        # preferir tabulación y punto y coma (la coma puede ser decimal)
        for sep in ("\t", ";", ",", "|"):
            if sep in candidatos:
                return sep
    return "espacios"


def _convertir(s, coma_decimal):
    if s is None:
        return None
    s = s.strip()
    if s == "":
        return None
    if _CERO_INICIAL.match(s):
        return s
    if _NUM.match(s):
        try:
            x = float(s)
            return int(x) if x.is_integer() and abs(x) < 1e15 and "." not in s and "e" not in s.lower() else x
        except ValueError:
            return s
    if coma_decimal:
        if _NUM_COMA.match(s):
            return float(s.replace(",", "."))
        if _MILES_PUNTO.match(s):
            return float(s.replace(".", "").replace(",", "."))
    return s


def _partir_espacios(linea):
    """Separa por espacios respetando comillas dobles (formato de Pedigree Viewer)."""
    campos = []
    actual = []
    en_comillas = False
    tuvo_comillas = False
    for ch in linea:
        if ch == '"':
            en_comillas = not en_comillas
            tuvo_comillas = True
            continue
        if ch in " \t" and not en_comillas:
            if actual or tuvo_comillas:
                campos.append("".join(actual))
                actual = []
                tuvo_comillas = False
            continue
        actual.append(ch)
    if actual or tuvo_comillas:
        campos.append("".join(actual))
    return campos


def _leer_texto(datos, separador=None):
    texto = _decodificar(datos)
    sep = separador or detectar_separador(texto)
    if sep == "espacios":
        filas_txt = [_partir_espacios(l) for l in texto.splitlines() if l.strip()]
    else:
        filas_txt = list(csv.reader(io.StringIO(texto), delimiter=sep))
    coma_decimal = sep in (";", "\t", "espacios", "|")
    if sep == ",":
        coma_decimal = False
    filas = [[_convertir(c, coma_decimal) for c in f] for f in filas_txt]
    if sum(len(f) for f in filas) > LIMITE_CELDAS:
        raise ErrorArchivo("El archivo es demasiado grande (más de 12 millones de celdas).")
    nombre_sep = {",": "coma", ";": "punto y coma", "\t": "tabulación", "|": "barra", "espacios": "espacios"}[sep]
    return [{"nombre": "Hoja1", "filas": _recortar(filas)}], nombre_sep


def leer(datos, nombre, separador=None):
    ext = ("." + nombre.rsplit(".", 1)[-1].lower()) if "." in nombre else ""
    info = {}
    if ext == ".ods":
        raise ErrorArchivo("Los archivos .ods no son compatibles: guárdelo como .xlsx o .csv desde LibreOffice.")
    if ext not in EXT_EXCEL and ext not in EXT_TEXTO and ext not in (".xls", ""):
        raise ErrorArchivo(f"Formato no reconocido ({ext}). Use .xlsx, .xls, .csv, .txt, .ped o .dat.")
    # el contenido manda sobre la extensión: muchos sistemas exportan CSV o HTML con extensión .xls,
    # y a veces un .xlsx se renombra como .xls
    if datos[:4] == b"PK\x03\x04":
        hojas = _leer_xlsx(datos)
    elif datos[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        hojas = _leer_xls(datos)
    elif datos.lstrip()[:1] == b"<" and ext in (".xls", ".xlsx"):
        raise ErrorArchivo("El archivo es una página web guardada con extensión de Excel. Ábralo en Excel y "
                           "guárdelo como libro de Excel (.xlsx) o como CSV.")
    elif ext in EXT_EXCEL:
        raise ErrorArchivo("El archivo no es un libro de Excel válido (puede estar dañado o protegido con contraseña).")
    else:
        hojas, info["separador"] = _leer_texto(datos, separador)
        base = nombre.rsplit(".", 1)[0] if "." in nombre else nombre
        hojas[0]["nombre"] = _nombre_hoja(base, set())
    hojas = [h for h in hojas if h["filas"]] or [{"nombre": "Hoja1", "filas": []}]
    return {"nombre": nombre, "hojas": hojas, **info}


# ================================================================ escritura
def _nombre_hoja(nombre, usados):
    limpio = re.sub(r"[\[\]\:\*\?\/\\]", "-", str(nombre or "Hoja")).strip()[:31] or "Hoja"
    base, k = limpio, 2
    while limpio.lower() in usados:
        sufijo = f" ({k})"
        limpio = base[:31 - len(sufijo)] + sufijo
        k += 1
    usados.add(limpio.lower())
    return limpio


def _estilos():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {
        "cab_fill": PatternFill("solid", fgColor="101010"),
        "cab_font": Font(bold=True, color="FFFFFF"),
        "titulo": Font(bold=True, size=13, color="101010"),
        "sub": Font(italic=True, color="666666"),
        "bold": Font(bold=True),
        "borde": Border(bottom=Side(style="thin", color="D9D9D9")),
        "centro": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "nota": Font(italic=True, size=9, color="666666"),
    }


def _ancho_columnas(ws, minimo=8, maximo=48):
    from openpyxl.utils import get_column_letter
    anchos = {}
    for fila in ws.iter_rows():
        for c in fila:
            if c.value is not None:
                largo = len(str(c.value)) if not isinstance(c.value, float) else 10
                anchos[c.column] = max(anchos.get(c.column, 0), largo)
    for col, largo in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimo, min(maximo, largo + 2))


def escribir_libro(libro):
    """libro = {"hojas": [{"nombre", "columnas": [..], "filas": [[..]]}]} -> bytes .xlsx"""
    from openpyxl import Workbook
    est = _estilos()
    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    for h in libro.get("hojas") or []:
        ws = wb.create_sheet(_nombre_hoja(h.get("nombre"), usados))
        cols = h.get("columnas") or []
        if cols:
            _agregar(ws, cols)
            for c in ws[1]:
                c.fill, c.font, c.alignment = est["cab_fill"], est["cab_font"], est["centro"]
            ws.freeze_panes = "A2"
        for fila in h.get("filas") or []:
            _agregar(ws, fila)
        _ancho_columnas(ws)
        if cols:
            ws.auto_filter.ref = ws.dimensions
    if not wb.worksheets:
        wb.create_sheet("Hoja1")
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _para_excel(v):
    if isinstance(v, (list, dict)):
        v = str(v)
    if isinstance(v, str):
        v = _ILEGALES.sub("", v)
        if len(v) > 32000:
            v = v[:32000]
    return v


def _agregar(ws, valores):
    """Agrega una fila; los textos que empiezan con «=» quedan como texto, no como fórmula."""
    fila = [_para_excel(v) for v in valores]
    ws.append(fila)
    if any(isinstance(v, str) and v.startswith("=") for v in fila):
        for c in ws[ws.max_row]:
            if c.data_type == "f":
                c.data_type = "s"


def escribir_resultados(resultados):
    """Exporta uno o varios resultados de análisis a un libro con formato."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    est = _estilos()
    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    for res in resultados:
        ws = wb.create_sheet(_nombre_hoja(res.get("titulo_corto") or res.get("titulo") or "Resultado", usados))
        ws.append([res.get("titulo") or "Resultado"])
        ws["A1"].font = est["titulo"]
        ws.append([res.get("subtitulo") or ""])
        ws["A2"].font = est["sub"]
        ws.append([f"JaaZ · {res.get('fecha', '')}"])
        ws["A3"].font = est["nota"]
        ws.append([])
        if res.get("indicadores"):
            _bloque(ws, est, "Resultados principales",
                    ["Parámetro", "Estimación", "Error estándar", "IC inferior", "IC superior", "Interpretación"],
                    [[i.get("etiqueta"), i.get("valor"), i.get("ee"),
                      (i.get("ic") or [None, None])[0], (i.get("ic") or [None, None])[1],
                      " · ".join(x for x in [i.get("nivel"), i.get("descripcion")] if x)]
                     for i in res["indicadores"]])
        for t in res.get("tablas") or []:
            cols = t.get("columnas") or []
            filas = t.get("filas_completas") or t.get("filas") or []
            _bloque(ws, est, t.get("titulo"), [c.get("titulo") for c in cols],
                    [[f.get(c["clave"]) for c in cols] for f in filas], t.get("nota"))
        if res.get("formulas"):
            _bloque(ws, est, "Fórmulas", ["Concepto", "Expresión"],
                    [[f.get("titulo"), f.get("expresion")] for f in res["formulas"]])
        if res.get("notas"):
            _bloque(ws, est, "Notas", ["Tipo", "Nota"],
                    [[n.get("tipo", "").capitalize(), n.get("texto")] for n in res["notas"]])
        _ancho_columnas(ws, maximo=60)
        ws.column_dimensions[get_column_letter(1)].width = max(ws.column_dimensions["A"].width or 0, 28)
        hn = res.get("hoja_nueva")
        if hn and hn.get("filas"):
            wh = wb.create_sheet(_nombre_hoja(hn.get("nombre") or "Datos", usados))
            _agregar(wh, hn.get("columnas") or [])
            for c in wh[1]:
                c.fill, c.font, c.alignment = est["cab_fill"], est["cab_font"], est["centro"]
            for fila in hn["filas"]:
                _agregar(wh, fila)
            wh.freeze_panes = "A2"
            _ancho_columnas(wh)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _bloque(ws, est, titulo, cabecera, filas, nota=None):
    _agregar(ws, [titulo])
    ws.cell(ws.max_row, 1).font = est["bold"]
    _agregar(ws, cabecera)
    for c in ws[ws.max_row]:
        c.fill, c.font, c.alignment = est["cab_fill"], est["cab_font"], est["centro"]
    for f in filas:
        _agregar(ws, f)
        for c in ws[ws.max_row]:
            if isinstance(c.value, float):
                c.number_format = "0.00E+00" if 0 < abs(c.value) < 1e-4 else "#,##0.0000"
    if nota:
        _agregar(ws, [nota])
        ws.cell(ws.max_row, 1).font = est["nota"]
    ws.append([])
